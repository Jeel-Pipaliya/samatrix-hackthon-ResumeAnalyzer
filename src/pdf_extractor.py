import io
import re
from typing import Dict, Any, List, Optional, Tuple
from PIL import Image

import pypdf
import pdfplumber
import pypdfium2

try:
    import winocr
    HAS_WINOCR = True
except Exception:
    HAS_WINOCR = False

try:
    import docx
    HAS_DOCX = True
except Exception:
    HAS_DOCX = False


class PDFExtractor:
    """
    Enterprise-grade, fault-tolerant resume extractor supporting:
    - PDF (Multi-engine waterfall: pdfplumber -> pypdfium2 -> pypdf -> Windows Native OCR)
    - Images (PNG, JPG, JPEG via Windows Native OCR)
    - Word Documents (DOCX)
    - Raw text and streams
    """

    SECTION_PATTERNS = {
        "Executive Summary": r"(summary|professional summary|executive summary|career overview|profile|about me)",
        "Work Experience": r"(experience|work experience|employment history|work history|professional experience)",
        "Education": r"(education|academic background|qualifications|academic history|degrees)",
        "Skills & Competencies": r"(skills|technical skills|core competencies|expertise|technologies)",
        "Certifications": r"(certifications|certificates|licenses|credentials)",
        "Projects": r"(projects|key projects|academic projects|personal projects)",
        "Languages": r"(languages|language proficiency)",
        "Awards & Honors": r"(awards|honors|achievements|recognition)"
    }

    @classmethod
    def _extract_from_pdf_plumber(cls, raw_bytes: bytes) -> str:
        """Engine 1: pdfplumber"""
        text_parts = []
        try:
            with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                for page in pdf.pages:
                    txt = page.extract_text(layout=False) or page.extract_text()
                    if txt and txt.strip():
                        text_parts.append(txt.strip())
        except Exception:
            pass
        return "\n\n".join(text_parts).strip()

    @classmethod
    def _extract_from_pdfium(cls, raw_bytes: bytes) -> str:
        """Engine 2: Chromium PDFium via pypdfium2"""
        text_parts = []
        try:
            doc = pypdfium2.PdfDocument(io.BytesIO(raw_bytes))
            for page in doc:
                textpage = page.get_textpage()
                txt = textpage.get_text_range()
                if txt and txt.strip():
                    text_parts.append(txt.strip())
        except Exception:
            pass
        return "\n\n".join(text_parts).strip()

    @classmethod
    def _extract_from_pypdf(cls, raw_bytes: bytes) -> str:
        """Engine 3: pypdf"""
        text_parts = []
        try:
            reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
            for page in reader.pages:
                txt = page.extract_text()
                if txt and txt.strip():
                    text_parts.append(txt.strip())
        except Exception:
            pass
        return "\n\n".join(text_parts).strip()

    @classmethod
    def _extract_from_ocr(cls, raw_bytes: bytes) -> Tuple[str, int]:
        """Engine 4: Windows Native OCR (winocr) on rendered PDF pages"""
        if not HAS_WINOCR:
            return "", 0

        text_parts = []
        page_count = 0
        try:
            doc = pypdfium2.PdfDocument(io.BytesIO(raw_bytes))
            page_count = len(doc)
            for page in doc:
                # Render page at 2x resolution for sharp text recognition
                pil_img = page.render(scale=2).to_pil()
                res = winocr.recognize_pil_sync(pil_img, 'en')
                if res and "text" in res and res["text"].strip():
                    text_parts.append(res["text"].strip())
        except Exception:
            pass
        return "\n\n".join(text_parts).strip(), page_count

    @classmethod
    def _extract_from_docx(cls, raw_bytes: bytes) -> str:
        """Word document extraction"""
        if not HAS_DOCX:
            return ""
        text_parts = []
        try:
            doc = docx.Document(io.BytesIO(raw_bytes))
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
            for t in doc.tables:
                for row in t.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))
        except Exception:
            pass
        return "\n\n".join(text_parts).strip()

    @classmethod
    def _extract_from_image(cls, raw_bytes: bytes) -> str:
        """Direct OCR on image uploads (PNG, JPG)"""
        if not HAS_WINOCR:
            return ""
        try:
            img = Image.open(io.BytesIO(raw_bytes))
            res = winocr.recognize_pil_sync(img, 'en')
            return res.get("text", "").strip()
        except Exception:
            return ""

    @classmethod
    def extract_text_and_meta(cls, file_source: Any, file_name: str = "") -> Dict[str, Any]:
        """
        Extracts raw text, pages, word count, contact details, and detected sections
        using the fault-tolerant Multi-Engine Waterfall.
        """
        # 1. Resolve raw bytes
        if hasattr(file_source, "getvalue"):
            raw_bytes = file_source.getvalue()
        elif isinstance(file_source, bytes):
            raw_bytes = file_source
        elif isinstance(file_source, str):
            with open(file_source, "rb") as f:
                raw_bytes = f.read()
            if not file_name:
                file_name = file_source
        else:
            raw_bytes = bytes(file_source)

        text = ""
        page_count = 1
        extraction_method = "Unknown"
        file_lower = file_name.lower()

        # Handle DOCX
        if file_lower.endswith(".docx"):
            text = cls._extract_from_docx(raw_bytes)
            extraction_method = "python-docx"

        # Handle Images
        elif file_lower.endswith((".png", ".jpg", ".jpeg", ".bmp", ".webp")):
            text = cls._extract_from_image(raw_bytes)
            extraction_method = "Windows Native OCR (winocr)"

        # Handle Plain Text
        elif file_lower.endswith(".txt"):
            try:
                text = raw_bytes.decode("utf-8", errors="ignore")
                extraction_method = "Plain Text UTF-8"
            except Exception:
                text = str(raw_bytes)

        # Handle PDF (Waterfall Architecture)
        else:
            # Step 1: Try pdfplumber
            text = cls._extract_from_pdf_plumber(raw_bytes)
            if len(text.strip()) >= 30:
                extraction_method = "pdfplumber (Layout Parser)"
                try:
                    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                        page_count = len(pdf.pages)
                except Exception:
                    page_count = 1

            # Step 2: Try pypdfium2 (Chrome Engine)
            if len(text.strip()) < 30:
                pdfium_text = cls._extract_from_pdfium(raw_bytes)
                if len(pdfium_text.strip()) >= 30:
                    text = pdfium_text
                    extraction_method = "pypdfium2 (Chromium Engine)"
                    try:
                        doc = pypdfium2.PdfDocument(io.BytesIO(raw_bytes))
                        page_count = len(doc)
                    except Exception:
                        page_count = 1

            # Step 3: Try pypdf
            if len(text.strip()) < 30:
                pypdf_text = cls._extract_from_pypdf(raw_bytes)
                if len(pypdf_text.strip()) >= 30:
                    text = pypdf_text
                    extraction_method = "pypdf (Stream Parser)"
                    try:
                        reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
                        page_count = len(reader.pages)
                    except Exception:
                        page_count = 1

            # Step 4: Fallback to Native Windows OCR (winocr) for scanned / flat PDFs
            if len(text.strip()) < 30 and HAS_WINOCR:
                ocr_text, ocr_pages = cls._extract_from_ocr(raw_bytes)
                if len(ocr_text.strip()) >= 20:
                    text = ocr_text
                    extraction_method = "Windows Native OCR (winocr)"
                    page_count = max(ocr_pages, 1)

        text = text.strip()
        words = text.split()
        word_count = len(words)
        char_count = len(text)
        reading_time_min = round(word_count / 200, 1) if word_count > 0 else 0.5

        # Extract contact information
        emails = list(set(re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)))
        phones = list(set(re.findall(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)))
        urls = list(set(re.findall(r'(https?://[^\s]+|www\.[^\s]+|linkedin\.com/in/[^\s]+|github\.com/[^\s]+)', text)))

        # Detect resume sections
        detected_sections = {}
        for section, pattern in cls.SECTION_PATTERNS.items():
            matches = list(re.finditer(rf'(?im)^[ \t]*({pattern})[ \t]*[:\-]?$', text))
            if not matches:
                matches = list(re.finditer(rf'(?i)\b({pattern})\b', text))
            detected_sections[section] = len(matches) > 0

        # Estimate structure completeness score (0 to 100)
        score = 25  # baseline
        if word_count >= 250:
            score += 25
        elif word_count >= 100:
            score += 15

        if detected_sections.get("Work Experience"):
            score += 15
        if detected_sections.get("Education"):
            score += 15
        if detected_sections.get("Skills & Competencies"):
            score += 10
        if emails or phones:
            score += 10

        return {
            "text": text,
            "page_count": max(page_count, 1),
            "word_count": word_count,
            "char_count": char_count,
            "reading_time_minutes": reading_time_min,
            "extraction_method": extraction_method,
            "contacts": {
                "emails": emails,
                "phones": phones,
                "links": urls
            },
            "detected_sections": detected_sections,
            "resume_health_score": min(score, 100)
        }
